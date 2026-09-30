"""
CoalMind focused LLM-assisted report generator.

Flow:
1. User describes what the report should be about.
2. Gemini converts the request into a structured report plan.
3. SQLite is searched only for evidence relevant to that plan.
4. Gemini analyzes only the retrieved evidence.
5. Python creates a focused Word document using python-docx.

Important:
- SQLite remains the source of truth for numbers.
- Gemini does not get the entire uploaded document collection.
- Gemini is instructed not to invent facts or numbers.
- Every generated finding should reference evidence IDs.
"""

import datetime
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from . import db, llm


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _h(doc, text, level=1):
    return doc.add_heading(str(text), level=level)


def _cell_shade(cell, hex_color="1F497D"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)

    tcPr.append(shd)


def make_table(doc, columns, rows):
    """Create a simple readable Word table."""

    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Light List Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header = table.rows[0].cells

    for i, col in enumerate(columns):
        header[i].text = str(col)

        if header[i].paragraphs and header[i].paragraphs[0].runs:
            run = header[i].paragraphs[0].runs[0]
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        _cell_shade(header[i])

    for row in rows:
        cells = table.add_row().cells

        for i, value in enumerate(row):
            cells[i].text = "" if value is None else str(value)

    return table


def _clean_text(value, max_chars=700):
    if value is None:
        return ""

    text = str(value)
    text = re.sub(r"\s+", " ", text).strip()

    if len(text) > max_chars:
        text = text[:max_chars] + "..."

    return text


def _normalise_list(value):
    if value is None:
        return []

    if isinstance(value, list):
        return [
            str(x).strip()
            for x in value
            if str(x).strip()
        ]

    if isinstance(value, str):
        return [
            x.strip()
            for x in re.split(r"[,;\n]", value)
            if x.strip()
        ]

    return [str(value)]


# ---------------------------------------------------------------------------
# 1. Report planning
# ---------------------------------------------------------------------------

PLAN_SYSTEM = """
You are the report planning assistant for CoalMind, an AI system used for
Indian coal-sector administrative and management reporting.

Your task is to understand what the user wants the report to be ABOUT.

Do NOT answer the report request yet.

Create a focused report plan.

Return ONLY valid JSON in this structure:

{
  "title": "short report title",
  "purpose": "what this report is intended to explain",
  "focus_topics": ["topic 1", "topic 2"],
  "keywords": ["keyword1", "keyword2"],
  "entities": ["CIL", "CMPDI"],
  "periods": ["2024-25", "2023-24"],
  "sections": [
    "Executive Summary",
    "Key Figures",
    "Analysis",
    "Key Observations",
    "Data Quality"
  ],
  "exclude": [
    "unrelated information"
  ]
}

Rules:
- Keep the report focused on the user's request.
- Do not add unrelated Ministry/coal-sector topics.
- Extract years/financial years when explicitly present.
- Extract organizations/entities when explicitly relevant.
- Use 5-15 useful keywords.
- Use 3-7 report sections.
- If the user did not specify a year, do not invent one.
- If the user did not specify an organization, do not invent one.
"""


def _fallback_plan(request):
    """
    Deterministic fallback when Gemini is unavailable.
    """

    text = request.strip()

    years = re.findall(r"\b20\d{2}[-–/]\d{2}\b", text)
    years = [
        y.replace("–", "-").replace("/", "-")
        for y in years
    ]

    known_entities = [
        "CIL",
        "CMPDI",
        "BCCL",
        "CCL",
        "ECL",
        "MCL",
        "NCL",
        "SECL",
        "WCL",
        "NEC",
        "SCCL",
    ]

    entities = []

    for entity in known_entities:
        if re.search(rf"\b{re.escape(entity)}\b", text, re.I):
            entities.append(entity)

    words = re.findall(r"[A-Za-z][A-Za-z-]{3,}", text.lower())

    stop_words = {
        "generate",
        "prepare",
        "create",
        "report",
        "about",
        "with",
        "include",
        "only",
        "based",
        "using",
        "from",
        "uploaded",
        "documents",
        "information",
        "should",
        "during",
        "between",
        "please",
        "make",
        "give",
        "show",
    }

    keywords = []

    for word in words:
        if word not in stop_words and word not in keywords:
            keywords.append(word)

    return {
        "title": "CoalMind Focused Report",
        "purpose": text,
        "focus_topics": keywords[:6],
        "keywords": keywords[:12],
        "entities": entities,
        "periods": years,
        "sections": [
            "Executive Summary",
            "Key Figures",
            "Analysis",
            "Key Observations",
            "Data Quality",
        ],
        "exclude": [
            "Information unrelated to the requested report topic"
        ],
    }


def build_report_plan(request):
    """
    Ask Gemini to understand the user's reporting objective.

    Returns:
        dict report plan
    """

    request = (request or "").strip()

    if not request:
        return _fallback_plan("General coal-sector report")

    response = llm.chat(
        PLAN_SYSTEM,
        f"User report request:\n{request}",
        json_mode=True,
        max_tokens=1000,
    )

    if response:
        plan = llm.extract_json(response)

        if isinstance(plan, dict):
            plan.setdefault("title", "CoalMind Focused Report")
            plan.setdefault("purpose", request)
            plan.setdefault("focus_topics", [])
            plan.setdefault("keywords", [])
            plan.setdefault("entities", [])
            plan.setdefault("periods", [])
            plan.setdefault(
                "sections",
                [
                    "Executive Summary",
                    "Key Figures",
                    "Analysis",
                    "Key Observations",
                    "Data Quality",
                ],
            )
            plan.setdefault("exclude", [])

            plan["focus_topics"] = _normalise_list(plan["focus_topics"])
            plan["keywords"] = _normalise_list(plan["keywords"])
            plan["entities"] = _normalise_list(plan["entities"])
            plan["periods"] = _normalise_list(plan["periods"])
            plan["sections"] = _normalise_list(plan["sections"])
            plan["exclude"] = _normalise_list(plan["exclude"])

            return plan

    return _fallback_plan(request)


# ---------------------------------------------------------------------------
# 2. Evidence retrieval
# ---------------------------------------------------------------------------

def _build_like_condition(columns, terms):
    """
    Build:

    (
      LOWER(column1) LIKE ?
      OR LOWER(column2) LIKE ?
      ...
    )

    for each keyword.
    """

    conditions = []
    params = []

    for term in terms:
        term = str(term).strip().lower()

        if not term:
            continue

        for column in columns:
            conditions.append(f"LOWER(COALESCE({column}, '')) LIKE ?")
            params.append(f"%{term}%")

    if not conditions:
        return "1=0", []

    return "(" + " OR ".join(conditions) + ")", params


def _retrieve_evidence(plan, max_facts=70, max_chunks=15):
    """
    Retrieve only evidence related to the report plan.

    Returns:
        {
            "facts": [...],
            "chunks": [...],
            "flags": [...],
            "documents": [...]
        }
    """

    con = db.connect()

    keywords = list(plan.get("keywords", []))
    topics = list(plan.get("focus_topics", []))
    entities = list(plan.get("entities", []))
    periods = list(plan.get("periods", []))

    terms = []

    for value in keywords + topics + entities:
        value = str(value).strip()

        if value and value.lower() not in [x.lower() for x in terms]:
            terms.append(value)

    # Keep the search manageable.
    terms = terms[:18]

    # -------------------------------------------------------
    # Facts
    # -------------------------------------------------------

    fact_columns = [
        "f.entity",
        "f.section",
        "f.column_label",
        "f.qualifier",
        "f.raw_text",
        "t.title",
        "t.sheet",
    ]

    like_sql, like_params = _build_like_condition(
        fact_columns,
        terms,
    )

    period_condition = ""
    period_params = []

    if periods:
        placeholders = ",".join("?" for _ in periods)

        period_condition = f"""
            AND (
                f.period IN ({placeholders})
                OR f.column_label IN ({placeholders})
            )
        """

        period_params = periods + periods

    sql = f"""
        SELECT
            f.id,
            f.entity,
            f.section,
            f.column_label,
            f.period,
            f.value,
            f.unit,
            f.raw_text,
            f.cell,
            f.sheet,
            f.confidence,
            f.status,
            t.title AS table_title,
            d.filename
        FROM facts f
        LEFT JOIN src_tables t ON t.id = f.table_id
        LEFT JOIN documents d ON d.id = f.doc_id
        WHERE {like_sql}
        {period_condition}
        ORDER BY
            CASE
                WHEN f.status = 'flagged' THEN 0
                WHEN f.confidence = 'high' THEN 1
                WHEN f.confidence = 'medium' THEN 2
                ELSE 3
            END,
            f.id DESC
        LIMIT ?
    """

    try:
        fact_rows = con.execute(
            sql,
            like_params + period_params + [max_facts],
        ).fetchall()
    except Exception:
        fact_rows = []

    facts = []

    for row in fact_rows:
        facts.append(
            {
                "evidence_id": f"F{len(facts) + 1}",
                "id": row["id"],
                "entity": row["entity"],
                "section": row["section"],
                "column_label": row["column_label"],
                "period": row["period"],
                "value": row["value"],
                "unit": row["unit"],
                "raw_text": _clean_text(row["raw_text"]),
                "cell": row["cell"],
                "sheet": row["sheet"],
                "confidence": row["confidence"],
                "status": row["status"],
                "table_title": row["table_title"],
                "filename": row["filename"],
            }
        )

    # -------------------------------------------------------
    # Text chunks
    # -------------------------------------------------------

    chunk_conditions = []
    chunk_params = []

    for term in terms[:12]:
        chunk_conditions.append("LOWER(text) LIKE ?")
        chunk_params.append(f"%{term.lower()}%")

    chunks = []

    if chunk_conditions:
        chunk_sql = f"""
            SELECT
                c.id,
                c.doc_id,
                c.loc,
                c.kind,
                c.text,
                d.filename
            FROM chunks c
            LEFT JOIN documents d ON d.id = c.doc_id
            WHERE {" OR ".join(chunk_conditions)}
            ORDER BY c.id DESC
            LIMIT ?
        """

        try:
            chunk_rows = con.execute(
                chunk_sql,
                chunk_params + [max_chunks],
            ).fetchall()
        except Exception:
            chunk_rows = []

        for row in chunk_rows:
            chunks.append(
                {
                    "evidence_id": f"C{len(chunks) + 1}",
                    "id": row["id"],
                    "doc_id": row["doc_id"],
                    "loc": row["loc"],
                    "kind": row["kind"],
                    "text": _clean_text(row["text"], 1100),
                    "filename": row["filename"],
                }
            )

    # -------------------------------------------------------
    # Relevant validation flags
    # -------------------------------------------------------

    flags = []

    fact_ids = [x["id"] for x in facts]

    if fact_ids:
        placeholders = ",".join("?" for _ in fact_ids)

        flag_rows = con.execute(
            f"""
            SELECT
                fl.id,
                fl.fact_id,
                fl.code,
                fl.severity,
                fl.detail,
                fl.resolved
            FROM flags fl
            WHERE fl.resolved = 0
              AND fl.fact_id IN ({placeholders})
            ORDER BY
                CASE fl.severity
                    WHEN 'error' THEN 0
                    WHEN 'warn' THEN 1
                    ELSE 2
                END,
                fl.id DESC
            LIMIT 30
            """,
            fact_ids,
        ).fetchall()

        for row in flag_rows:
            flags.append(
                {
                    "evidence_id": f"V{len(flags) + 1}",
                    "fact_id": row["fact_id"],
                    "code": row["code"],
                    "severity": row["severity"],
                    "detail": _clean_text(row["detail"], 500),
                }
            )

    # -------------------------------------------------------
    # Documents represented by evidence
    # -------------------------------------------------------

    filenames = []

    for item in facts:
        if item["filename"] and item["filename"] not in filenames:
            filenames.append(item["filename"])

    for item in chunks:
        if item["filename"] and item["filename"] not in filenames:
            filenames.append(item["filename"])

    con.close()

    return {
        "facts": facts,
        "chunks": chunks,
        "flags": flags,
        "documents": filenames,
    }


# ---------------------------------------------------------------------------
# 3. Prepare evidence for Gemini
# ---------------------------------------------------------------------------

def _evidence_prompt(evidence):
    blocks = []

    for fact in evidence["facts"]:
        blocks.append(
            f"""
[{fact['evidence_id']}]
TYPE: STRUCTURED FACT
Source file: {fact['filename']}
Sheet: {fact['sheet']}
Cell: {fact['cell']}
Table: {fact['table_title']}
Entity: {fact['entity']}
Section: {fact['section']}
Period: {fact['period']}
Column: {fact['column_label']}
Value: {fact['value']}
Unit: {fact['unit']}
Confidence: {fact['confidence']}
Status: {fact['status']}
Raw text: {fact['raw_text']}
""".strip()
        )

    for chunk in evidence["chunks"]:
        blocks.append(
            f"""
[{chunk['evidence_id']}]
TYPE: DOCUMENT TEXT
Source file: {chunk['filename']}
Location: {chunk['loc']}
Text:
{chunk['text']}
""".strip()
        )

    for flag in evidence["flags"]:
        blocks.append(
            f"""
[{flag['evidence_id']}]
TYPE: VALIDATION FINDING
Code: {flag['code']}
Severity: {flag['severity']}
Detail: {flag['detail']}
Related fact ID: {flag['fact_id']}
""".strip()
        )

    if not blocks:
        return "NO RELEVANT EVIDENCE WAS FOUND."

    return "\n\n---\n\n".join(blocks)


# ---------------------------------------------------------------------------
# 4. LLM report analysis
# ---------------------------------------------------------------------------

REPORT_SYSTEM = """
You are the report-analysis assistant for CoalMind.

You receive:
1. A user's report objective.
2. A report plan.
3. Evidence retrieved from CoalMind's SQLite database.

Generate a focused report based ONLY on the supplied evidence.

IMPORTANT RULES:

- Do NOT reproduce the uploaded documents.
- Do NOT summarize every uploaded document.
- Use only information relevant to the user's requested topic.
- Do NOT invent numbers, dates, organizations, causes, achievements or conclusions.
- Every numerical claim MUST have an evidence ID such as [F3].
- Every source-based statement should have an evidence ID such as [C2].
- Validation findings must cite their [V#] evidence ID.
- If evidence is missing, explicitly say:
  "The uploaded data does not provide sufficient evidence for this point."
- Do not assume that a missing value means zero.
- Explain figures in simple language so a non-technical administrative user can understand them.
- Explain what a trend/value means, but do not invent reasons for the trend.
- Keep the report focused on the requested objective.
- Respect the report sections supplied in the plan.
- Mention important limitations when the evidence is incomplete.

Return ONLY JSON:

{
  "executive_summary": "short focused summary with evidence IDs",
  "sections": [
    {
      "heading": "section heading",
      "explanation": "simple explanation of what the evidence means",
      "points": [
        "point with evidence ID",
        "point with evidence ID"
      ]
    }
  ],
  "key_findings": [
    "finding with evidence ID"
  ],
  "data_quality": [
    "relevant validation observation with evidence ID"
  ],
  "limitations": [
    "limitation"
  ]
}

Do not include markdown fences.
"""


def _generate_analysis(request, plan, evidence):
    evidence_text = _evidence_prompt(evidence)

    user_prompt = f"""
USER REPORT REQUEST:
{request}

REPORT PLAN:
{json.dumps(plan, indent=2)}

RETRIEVED EVIDENCE:
{evidence_text}
"""

    response = llm.chat(
        REPORT_SYSTEM,
        user_prompt,
        json_mode=True,
        max_tokens=2600,
    )

    if response:
        parsed = llm.extract_json(response)

        if isinstance(parsed, dict):
            parsed.setdefault("executive_summary", "")
            parsed.setdefault("sections", [])
            parsed.setdefault("key_findings", [])
            parsed.setdefault("data_quality", [])
            parsed.setdefault("limitations", [])

            return parsed

    return _fallback_analysis(plan, evidence)


def _fallback_analysis(plan, evidence):
    """
    Deterministic fallback.

    It does NOT try to invent a narrative. It simply reports the
    relevant evidence that was found.
    """

    points = []

    for fact in evidence["facts"][:15]:
        label = fact["entity"] or "Value"

        if fact["period"]:
            label += f" ({fact['period']})"

        points.append(
            f"{label}: {fact['value']} {fact['unit'] or ''} [{fact['evidence_id']}]"
        )

    summary = (
        "The report is limited to the evidence retrieved for the requested "
        "topic. An LLM-generated interpretation was not available."
    )

    sections = [
        {
            "heading": "Relevant Evidence",
            "explanation": (
                "The following figures and document excerpts were selected "
                "because they match the requested report topic."
            ),
            "points": points,
        }
    ]

    quality = []

    for flag in evidence["flags"]:
        quality.append(
            f"{flag['severity'].upper()} {flag['code']}: "
            f"{flag['detail']} [{flag['evidence_id']}]"
        )

    return {
        "executive_summary": summary,
        "sections": sections,
        "key_findings": points[:8],
        "data_quality": quality,
        "limitations": [
            "LLM analysis was unavailable, so the report contains a deterministic evidence summary."
        ],
    }


# ---------------------------------------------------------------------------
# 5. Citation resolution
# ---------------------------------------------------------------------------

def _add_citation_text(doc, text, evidence_map):
    """
    Adds text and makes evidence IDs visually distinct.

    Example:
       "Production increased [F3]."
    """

    if not text:
        return

    pattern = re.compile(r"\[(F|C|V)\d+\]")

    pos = 0

    for match in pattern.finditer(str(text)):
        before = str(text)[pos:match.start()]

        if before:
            doc.add_paragraph().add_run(before)

        citation = match.group(0)

        p = doc.add_paragraph()

        if citation in evidence_map:
            run = p.add_run(citation)
            run.bold = True
            run.font.color.rgb = RGBColor(31, 73, 125)

        else:
            p.add_run(citation)

        pos = match.end()

    remaining = str(text)[pos:]

    if remaining:
        doc.add_paragraph().add_run(remaining)


def _write_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(str(text))
    return p


# ---------------------------------------------------------------------------
# 6. Main report generator
# ---------------------------------------------------------------------------

def generate_report(out_path=None, report_request=None, plan=None):
    """
    Generate a focused Word report.

    Parameters:
        out_path:
            Output .docx path.

        report_request:
            Natural-language description of what the report should be about.

        plan:
            Optional already-created report plan.

    Returns:
        str path to generated DOCX.
    """

    if not report_request:
        report_request = (
            "Generate a focused report from the uploaded coal-sector data "
            "based on the most relevant available information."
        )

    out_path = Path(
        out_path
        or (db.get_path().parent / "CoalMind_Focused_Report.docx")
    )

    # -------------------------------------------------------
    # Step 1: Understand user objective
    # -------------------------------------------------------

    if plan is None:
        plan = build_report_plan(report_request)

    # -------------------------------------------------------
    # Step 2: Retrieve only relevant evidence
    # -------------------------------------------------------

    evidence = _retrieve_evidence(plan)

    # -------------------------------------------------------
    # Step 3: Gemini analyzes relevant evidence
    # -------------------------------------------------------

    analysis = _generate_analysis(
        report_request,
        plan,
        evidence,
    )

    # -------------------------------------------------------
    # Step 4: Build Word document
    # -------------------------------------------------------

    now = datetime.datetime.now().strftime("%d %b %Y %H:%M")

    doc = Document()

    # Basic styles
    for style_name in ("Normal", "Heading 1", "Heading 2"):
        try:
            doc.styles[style_name].font.name = "Calibri"
        except Exception:
            pass

    # -------------------------------------------------------
    # Cover
    # -------------------------------------------------------

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cover.add_run("\n\n\n")

    title_run = cover.add_run(
        plan.get("title", "CoalMind Focused Report")
    )

    title_run.bold = True
    title_run.font.size = Pt(28)
    title_run.font.color.rgb = RGBColor(31, 73, 125)

    subtitle = cover.add_run(
        "\n\nLLM-Assisted, Evidence-Grounded Administrative Report"
    )

    subtitle.font.size = Pt(14)

    generated = doc.add_paragraph(
        f"\nGenerated: {now}"
    )

    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # -------------------------------------------------------
    # 1. Report Scope
    # -------------------------------------------------------

    _h(doc, "1. Report Scope")

    p = doc.add_paragraph()
    p.add_run("Requested objective: ").bold = True
    p.add_run(str(report_request))

    p = doc.add_paragraph()
    p.add_run("Purpose: ").bold = True
    p.add_run(str(plan.get("purpose", "")))

    if plan.get("focus_topics"):
        p = doc.add_paragraph()
        p.add_run("Focus topics: ").bold = True
        p.add_run(", ".join(plan["focus_topics"]))

    if plan.get("entities"):
        p = doc.add_paragraph()
        p.add_run("Entities: ").bold = True
        p.add_run(", ".join(plan["entities"]))

    if plan.get("periods"):
        p = doc.add_paragraph()
        p.add_run("Periods: ").bold = True
        p.add_run(", ".join(plan["periods"]))

    if plan.get("exclude"):
        p = doc.add_paragraph()
        p.add_run("Excluded from focus: ").bold = True
        p.add_run(", ".join(plan["exclude"]))

    # -------------------------------------------------------
    # 2. Executive Summary
    # -------------------------------------------------------

    _h(doc, "2. Executive Summary")

    summary = analysis.get("executive_summary", "")

    if summary:
        _add_citation_text(doc, summary, {})
    else:
        doc.add_paragraph("No executive summary was generated.")

    # -------------------------------------------------------
    # 3. Relevant Evidence / Key Figures
    # -------------------------------------------------------

    _h(doc, "3. Relevant Evidence and Key Figures")

    if evidence["facts"]:
        rows = []

        for fact in evidence["facts"][:30]:
            rows.append(
                [
                    fact["evidence_id"],
                    fact["entity"],
                    fact["period"] or fact["column_label"],
                    fact["value"],
                    fact["unit"],
                    fact["filename"],
                ]
            )

        make_table(
            doc,
            [
                "ID",
                "Entity",
                "Period",
                "Value",
                "Unit",
                "Source",
            ],
            rows,
        )

    else:
        doc.add_paragraph(
            "No structured numerical evidence matching the requested topic was found."
        )

    # -------------------------------------------------------
    # 4. Explanation / Analysis
    # -------------------------------------------------------

    _h(doc, "4. Explanation and Analysis")

    sections = analysis.get("sections", [])

    if not sections:
        doc.add_paragraph(
            "No additional analysis was generated."
        )

    for section in sections:

        heading = section.get(
            "heading",
            "Analysis",
        )

        _h(doc, heading, level=2)

        explanation = section.get(
            "explanation",
            "",
        )

        if explanation:
            p = doc.add_paragraph()
            p.add_run("Explanation: ").bold = True
            p.add_run(str(explanation))

        points = section.get("points", [])

        for point in points:
            _write_bullet(doc, point)

    # -------------------------------------------------------
    # 5. Key Findings
    # -------------------------------------------------------

    _h(doc, "5. Key Findings")

    findings = analysis.get("key_findings", [])

    if findings:
        for finding in findings:
            _write_bullet(doc, finding)
    else:
        doc.add_paragraph(
            "No additional key findings were generated from the selected evidence."
        )

    # -------------------------------------------------------
    # 6. Data Quality
    # -------------------------------------------------------

    _h(doc, "6. Relevant Data Quality Findings")

    quality = analysis.get("data_quality", [])

    if quality:
        for item in quality:
            _write_bullet(doc, item)
    elif evidence["flags"]:
        for flag in evidence["flags"]:
            _write_bullet(
                doc,
                f"{flag['severity'].upper()} {flag['code']}: "
                f"{flag['detail']} [{flag['evidence_id']}]",
            )
    else:
        doc.add_paragraph(
            "No unresolved validation finding was associated with the selected evidence."
        )

    # -------------------------------------------------------
    # 7. Limitations
    # -------------------------------------------------------

    _h(doc, "7. Limitations and Interpretation Notes")

    limitations = analysis.get("limitations", [])

    if limitations:
        for limitation in limitations:
            _write_bullet(doc, limitation)
    else:
        doc.add_paragraph(
            "The report is limited to the data successfully ingested and "
            "retrieved for the requested topic."
        )

    # -------------------------------------------------------
    # 8. Sources
    # -------------------------------------------------------

    _h(doc, "8. Sources Used")

    if evidence["documents"]:
        for filename in evidence["documents"]:
            _write_bullet(doc, filename)
    else:
        doc.add_paragraph(
            "No source documents were matched to the requested topic."
        )

    # -------------------------------------------------------
    # 9. Evidence Reference
    # -------------------------------------------------------

    _h(doc, "9. Evidence Reference")

    evidence_map = {}

    for fact in evidence["facts"]:
        evidence_map[f"[{fact['evidence_id']}]"] = fact

        p = doc.add_paragraph()

        p.add_run(
            f"[{fact['evidence_id']}] "
        ).bold = True

        p.add_run(
            f"{fact['filename']} | "
            f"Sheet: {fact['sheet']} | "
            f"Cell: {fact['cell']} | "
            f"Entity: {fact['entity']} | "
            f"Period: {fact['period']} | "
            f"Value: {fact['value']} {fact['unit'] or ''}"
        )

    for chunk in evidence["chunks"]:
        evidence_map[f"[{chunk['evidence_id']}]"] = chunk

        p = doc.add_paragraph()

        p.add_run(
            f"[{chunk['evidence_id']}] "
        ).bold = True

        p.add_run(
            f"{chunk['filename']} | "
            f"Location: {chunk['loc']}"
        )

        p = doc.add_paragraph()
        p.add_run(chunk["text"])

    for flag in evidence["flags"]:
        evidence_map[f"[{flag['evidence_id']}]"] = flag

        p = doc.add_paragraph()

        p.add_run(
            f"[{flag['evidence_id']}] "
        ).bold = True

        p.add_run(
            f"{flag['severity'].upper()} | "
            f"{flag['code']} | "
            f"{flag['detail']}"
        )

    # -------------------------------------------------------
    # Footer
    # -------------------------------------------------------

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = footer.add_run(
            "CoalMind — Evidence-grounded report | "
            "Generated from ingested data"
        )

        run.font.size = Pt(8)

    # -------------------------------------------------------
    # Save
    # -------------------------------------------------------

    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc.save(str(out_path))

    db.audit(
        "REPORT",
        str(out_path),
        f"focused report generated; request={report_request[:200]}",
    )

    return str(out_path)